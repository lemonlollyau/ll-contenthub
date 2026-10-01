"""Builds templates/content-calendar-template.xlsx, matched to the app's importer.

Run: python3 scripts/build-calendar-template.py  (needs openpyxl)
Header names must stay in sync with src/lib/calendar/columns.ts.
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / "templates" / "content-calendar-template.xlsx"
FONT = "Arial"
LEMON = "F5B800"
INK = "1C1917"
REQUIRED_FILL = PatternFill("solid", start_color="FFF4CC")
EXAMPLE_FILL = PatternFill("solid", start_color="F5F5F4")
HEAD_FILL = PatternFill("solid", start_color=INK)
THIN = Side(style="thin", color="D6D3D1")
ROWS = 80  # rows pre-formatted for data entry

# (header, width, required, note)
SOCIAL = [
    ("#", 5, True, "Post number. Keep it unique and don't change it after importing: re-imports use it to update the same post."),
    ("Date", 12, True, "Real Excel date (e.g. 3/11/2026) or YYYY-MM-DD."),
    ("Time", 8, False, "Optional, 24-hour e.g. 18:30. Leave blank to use the client's posting-time rules."),
    ("Day", 6, False, "Optional, e.g. Tue. Only used as a hint for weekend times."),
    ("Phase", 12, False, "Optional grouping, e.g. Pre-sale, Sale, Post-sale."),
    ("Slot", 6, False, "AM = 8:00am. PM = the pillar's evening time. Blank = weekday/weekend time."),
    ("Platform", 14, True, "Instagram, Facebook, or both as 'Instagram / Facebook'. Also LinkedIn, TikTok, Pinterest, Threads."),
    ("Format", 10, True, "Feed, Carousel, Reel, Story, Square or Text."),
    ("Pillar", 16, False, "Must match a pillar name in the client's posting-time table to get its times."),
    ("Moment / Offer", 20, False, "The offer or moment this post supports. 'Ends midnight', 'last day' etc. move it to 8:30pm."),
    ("Hook", 28, False, "Opening line / idea. Used for image matching."),
    ("Caption (draft)", 50, True, "The full caption, exactly as it should go out. Blank lines separate paragraphs; single line breaks are kept."),
    ("Hashtags", 24, False, "Space-separated, e.g. #redlighttherapy #skincare. Added as the last paragraph."),
    ("CTA", 22, False, "Added as its own paragraph (before a compliance line) unless the caption already contains it."),
    ("First comment", 22, False, "Optional. Instagram, Facebook and LinkedIn only; posted by Buffer as the first comment."),
    ("Asset brief", 32, False, "Describe the image/video you want. For carousels say how many, e.g. '5 slides: ...'."),
    ("Asset group", 18, False, "Name of a subfolder in the client's Drive folder to choose from, e.g. 'Before and after'."),
    ("Asset file", 22, False, "Optional exact file name(s) or Drive link(s), comma-separated in carousel order. Overrides matching."),
    ("Asset status", 12, False, "Your tracking only."),
    ("Caption status", 12, False, "Set to Approved to lock the caption: the app will never edit it."),
    ("Owner", 10, False, "Your tracking only."),
    ("Scheduled", 10, False, "Your tracking only."),
    ("Live", 8, False, "Your tracking only."),
    ("Compliance note", 28, False, "Shown as a warning before pushing. Not sent to Buffer."),
]

SOCIAL_EXAMPLE = [
    1, "2026-11-03", "", "Tue", "Pre-sale", "AM", "Instagram / Facebook", "Carousel", "The Science Bit",
    "Black Friday early access",
    "Why red light works (in 60 seconds)",
    "Red light isn't magic, it's biology.\n\nWavelengths around 630–850nm reach your skin cells and support their natural energy production.\n\nSwipe for the 3-step routine →\n\nAlways read the label and follow the directions for use.",
    "#redlighttherapy #skinscience #fringeheals",
    "Save this for later",
    "",
    "Carousel, 4 slides: device close-up, face panel in use, routine steps graphic, product flat lay",
    "Product shots",
    "",
    "In Drive", "Approved", "Indy", "", "",
    "ARTG-listed device: no cure claims.",
]

EMAIL = [
    ("#", 5, True, "Email number. Keep it unique."),
    ("Date", 12, True, "Send date."),
    ("Time", 8, False, "Optional send time, 24-hour."),
    ("Moment / Offer", 20, False, "The offer or moment this email supports."),
    ("Subject line", 36, True, "Main subject line."),
    ("Preview text", 36, False, "Preview/preheader text."),
    ("Hook", 28, False, "Headline for the hero block."),
    ("Body copy", 50, False, "Draft body copy."),
    ("CTA", 20, False, "Button text."),
    ("Asset brief", 30, False, "Hero image you want."),
    ("Asset group", 18, False, "Drive subfolder to pick from."),
    ("Asset file", 22, False, "Optional exact file name or Drive link."),
    ("Owner", 10, False, "Your tracking only."),
    ("Compliance note", 28, False, "Shown as a warning."),
]

EMAIL_EXAMPLE = [
    1, "2026-11-27", "07:00", "Black Friday: 30% off", "Black Friday is here: 30% off every device",
    "Our biggest sale of the year ends Monday.", "Glow for less",
    "Our biggest sale of the year is live.\n\nEvery Fringe device is 30% off until midnight Monday.",
    "Shop the sale", "Face panel lifestyle shot, warm tones", "Lifestyle", "", "Indy",
    "Always read the label and follow the directions for use.",
]

LISTS = {
    "Platform options": ["Instagram", "Facebook", "Instagram / Facebook", "LinkedIn", "TikTok", "Pinterest", "Threads"],
    "Format options": ["Feed", "Carousel", "Reel", "Story", "Square", "Text"],
    "Slot options": ["AM", "PM"],
    "Asset status options": ["Needed", "Briefed", "In Drive", "Approved"],
    "Caption status options": ["Draft", "In review", "Approved"],
}


def font(**kw):
    return Font(name=FONT, **kw)


def build_sheet(ws, columns, example, list_ranges):
    ws.freeze_panes = "B2"
    ws.sheet_view.zoomScale = 110
    for i, (header, width, required, note) in enumerate(columns, start=1):
        c = ws.cell(row=1, column=i, value=header)
        c.font = font(bold=True, color=LEMON if required else "FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.comment = Comment(("REQUIRED. " if required else "") + note, "lemonlolly")
        c.comment.width, c.comment.height = 260, 110
        ws.column_dimensions[c.column_letter].width = width
    ws.row_dimensions[1].height = 30

    for r in range(2, ROWS + 2):
        for i, (header, _w, required, _n) in enumerate(columns, start=1):
            c = ws.cell(row=r, column=i)
            c.font = font(size=10)
            c.alignment = Alignment(vertical="top", wrap_text=header in ("Caption (draft)", "Body copy", "Asset brief", "Hook", "Compliance note"))
            c.border = Border(bottom=THIN)
            if header == "Date":
                c.number_format = "ddd d mmm yyyy"
            if header in ("Time", "#"):
                c.number_format = "@" if header == "Time" else "0"

    # Example row (row 2), greyed and labelled so it's obvious.
    from datetime import date
    for i, value in enumerate(example, start=1):
        c = ws.cell(row=2, column=i)
        header = columns[i - 1][0]
        if header == "Date" and value:
            y, m, d = map(int, value.split("-"))
            value = date(y, m, d)
        c.value = value if value != "" else None
        c.fill = EXAMPLE_FILL
        c.font = font(size=10, italic=True, color="57534E")
    ws.row_dimensions[2].height = 150

    headers = [c[0] for c in columns]
    for header, list_name in list_ranges.items():
        if header not in headers:
            continue
        col = ws.cell(row=1, column=headers.index(header) + 1).column_letter
        dv = DataValidation(type="list", formula1=f"'Lists'!${list_name}", allow_blank=True, showErrorMessage=False)
        dv.add(f"{col}2:{col}{ROWS + 1}")
        ws.add_data_validation(dv)


wb = Workbook()

readme = wb.active
readme.title = "Read me"
readme.column_dimensions["A"].width = 110
lines = [
    ("lemonlolly Content Hub: content calendar template", font(bold=True, size=16)),
    ("", None),
    ("How to use", font(bold=True, size=12)),
    ("1. Fill in the 'Schedule' tab (social posts) and, if needed, the 'Email' tab. One row per post.", None),
    ("2. Row 2 on each tab is a grey EXAMPLE. Delete it or overwrite it before importing, otherwise it will be imported.", None),
    ("3. Yellow headers are required: #, Date, Platform, Format and Caption (Schedule); #, Date and Subject line (Email).", None),
    ("4. Hover over any header to see what goes in it. Dropdowns are there for Platform, Format, Slot and status columns.", None),
    ("5. Don't rename the headers. You can reorder columns or add your own extra columns (the app ignores unknown ones).", None),
    ("6. In the app: Clients → the client → Calendar → Import calendar. Check the column matching, pick the month, Import.", None),
    ("", None),
    ("Rules the app applies", font(bold=True, size=12)),
    ("• Captions are used exactly as written. The CTA is added as its own paragraph (before a compliance line, if the caption ends with one), then the hashtags.", None),
    ("• Blank lines separate paragraphs. Single line breaks (price lists, quote stacks) are kept.", None),
    ("• Instagram captions over 2,200 characters, and placeholders in square brackets such as [DROP IN REAL REVIEW], are blocked from pushing.", None),
    ("• Every Instagram post needs an image or video; the app matches one from the client's Drive folder.", None),
    ("• No time? AM slot → 8:00, deadline posts (ends midnight, last day…) → 20:30, PM slot → the pillar's evening time, otherwise the pillar's weekday/weekend time.", None),
    ("• Re-importing the same month updates rows by their # number, keeping approved images and Buffer drafts. Keep # numbers stable.", None),
    ("• Carousel, Reel and Story formats go to Buffer via the API; the CSV fallback lists them for manual setup.", None),
    ("• A tab is treated as email if its name contains 'Email'. Other tabs with matching headers are treated as social posts.", None),
]
for i, (text, f) in enumerate(lines, start=1):
    c = readme.cell(row=i, column=1, value=text)
    c.font = f or font(size=11)
    c.alignment = Alignment(wrap_text=True, vertical="top")
readme.sheet_view.showGridLines = False

schedule = wb.create_sheet("Schedule")
email = wb.create_sheet("Email")

lists = wb.create_sheet("Lists")
ranges = {}
for col, (name, values) in enumerate(LISTS.items(), start=1):
    h = lists.cell(row=1, column=col, value=name)
    h.font = font(bold=True)
    for r, v in enumerate(values, start=2):
        lists.cell(row=r, column=col, value=v).font = font()
    letter = h.column_letter
    lists.column_dimensions[letter].width = 24
    ranges[name] = f"{letter}$2:${letter}${len(values) + 1}"
list_for = {
    "Platform": ranges["Platform options"],
    "Format": ranges["Format options"],
    "Slot": ranges["Slot options"],
    "Asset status": ranges["Asset status options"],
    "Caption status": ranges["Caption status options"],
}
build_sheet(schedule, SOCIAL, SOCIAL_EXAMPLE, list_for)
build_sheet(email, EMAIL, EMAIL_EXAMPLE, list_for)
schedule.sheet_properties.tabColor = LEMON
email.sheet_properties.tabColor = "A8A29E"

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(OUT)
